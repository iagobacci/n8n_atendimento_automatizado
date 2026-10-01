"""
=============================================================================
EXTRATOR DE PROFISSIONAIS DA SAÚDE - WHATSAPP COM COLUNAS SEPARADAS
=============================================================================
Gera automaticamente:
1. Arquivo XLSX (Excel nativo estilizado, colunas ajustadas e links clicáveis)
2. Arquivo CSV delimitado por ponto e vírgula (;) para abrir perfeitamente no Excel

Colunas geradas:
- Coluna A: Nome do Profissional / Clínica
- Coluna B: Especialidade
- Coluna C: DDD
- Coluna D: Número
- Coluna E: WhatsApp Completo
- Coluna F: Link WhatsApp (wa.me/55...)
- Coluna G: Cidade
- Coluna H: Estado (UF)
- Coluna I: Bairro
- Coluna J: CEP
- Coluna K: Endereço Completo
- Coluna L: Link da Fonte
=============================================================================
"""

import sys
import os
import re
import csv
import json
import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed

try:
    import requests
    from bs4 import BeautifulSoup
except ImportError:
    print("[ERRO] Execute: pip install requests beautifulsoup4")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36',
    'Accept-Language': 'pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7'
}

CIDADES_DISPONIVEIS = {
    '1': {'nome': 'São Paulo', 'uf': 'SP', 'slug': 'sao-paulo'},
    '2': {'nome': 'Rio de Janeiro', 'uf': 'RJ', 'slug': 'rio-de-janeiro'},
    '3': {'nome': 'Belo Horizonte', 'uf': 'MG', 'slug': 'belo-horizonte'},
    '4': {'nome': 'Curitiba', 'uf': 'PR', 'slug': 'curitiba'},
    '5': {'nome': 'Brasília', 'uf': 'DF', 'slug': 'brasilia'},
    '6': {'nome': 'Porto Alegre', 'uf': 'RS', 'slug': 'porto-alegre'},
    '7': {'nome': 'Campinas', 'uf': 'SP', 'slug': 'campinas'},
    '8': {'nome': 'Salvador', 'uf': 'BA', 'slug': 'salvador'},
    '9': {'nome': 'Fortaleza', 'uf': 'CE', 'slug': 'fortaleza'},
    '10': {'nome': 'Recife', 'uf': 'PE', 'slug': 'recife'},
    '11': {'nome': 'Goiânia', 'uf': 'GO', 'slug': 'goiania'},
    '12': {'nome': 'Florianópolis', 'uf': 'SC', 'slug': 'florianopolis'},
    '13': {'nome': 'Santos', 'uf': 'SP', 'slug': 'santos'},
    '14': {'nome': 'Ribeirão Preto', 'uf': 'SP', 'slug': 'ribeirao-preto'},
    '15': {'nome': 'Niterói', 'uf': 'RJ', 'slug': 'niteroi'},
    '16': {'nome': 'Santo André', 'uf': 'SP', 'slug': 'santo-andre'},
    '17': {'nome': 'São Bernardo do Campo', 'uf': 'SP', 'slug': 'sao-bernardo-do-campo'},
    '18': {'nome': 'Londrina', 'uf': 'PR', 'slug': 'londrina'}
}

def extract_mobile_and_wa(raw_tel, wa_tag_href, page_text):
    if wa_tag_href:
        m = re.search(r'wa\.me/(\d+)', wa_tag_href) or re.search(r'phone=(\d+)', wa_tag_href)
        if m:
            num = m.group(1)
            if num.startswith('55') and len(num) >= 12:
                num = num[2:]
            if len(num) in (10, 11):
                formatted = f"({num[:2]}) {num[2:7]}-{num[7:]}" if len(num) == 11 else f"({num[:2]}) {num[2:6]}-{num[6:]}"
                wa_url = f"https://wa.me/55{num}"
                return formatted, wa_url

    digits = re.sub(r'\D', '', raw_tel or '')
    if digits.startswith('55') and len(digits) >= 12:
        digits = digits[2:]
    
    if len(digits) == 11 and digits[2] == '9':
        formatted = f"({digits[:2]}) {digits[2:7]}-{digits[7:]}"
        wa_url = f"https://wa.me/55{digits}"
        return formatted, wa_url

    if page_text:
        mobiles = re.findall(r'\(?\b[1-9]{2}\)?\s*9\s*[6-9]\d{3}[-\s]?\d{4}\b', page_text)
        for mob in mobiles:
            m_digits = re.sub(r'\D', '', mob)
            if len(m_digits) == 11 and m_digits[2] == '9':
                formatted = f"({m_digits[:2]}) {m_digits[2:7]}-{m_digits[7:]}"
                wa_url = f"https://wa.me/55{m_digits}"
                return formatted, wa_url

    return None, None

def parse_address_parts(addr_str, city, uf):
    if not addr_str or addr_str == f"{city} - {uf}":
        return "", "", addr_str

    cep_match = re.search(r'\b\d{5}[-.]?\d{3}\b', addr_str)
    cep = cep_match.group(0) if cep_match else ""

    bairro = ""
    parts = [p.strip() for p in addr_str.split('-')]
    for p in parts:
        if city.lower() in p.lower() or uf.lower() in p.lower() or 'brasil' in p.lower():
            continue
        if len(p) > 2 and not any(char.isdigit() for char in p) and not any(w in p.lower() for w in ['rua', 'av', 'avenida', 'sala', 'loja', 'conj', 'andar']):
            bairro = p.strip()
            break

    return bairro, cep, addr_str

def scrape_dentmap_page(city_info, page_num):
    results = []
    slug = city_info['slug']
    url = f"https://dentmap.com.br/dentistas/{slug}?page={page_num}"
    try:
        r = requests.get(url, headers=HEADERS, timeout=12)
        if r.status_code != 200:
            return results
        soup = BeautifulSoup(r.text, 'html.parser')
        
        dentists_to_fetch = []
        for s in soup.find_all('script', type='application/ld+json'):
            try:
                data = json.loads(s.string)
                if isinstance(data, dict) and 'itemListElement' in data:
                    for el in data['itemListElement']:
                        it = el.get('item', {})
                        if it.get('@type') == 'Dentist' and it.get('url'):
                            dentists_to_fetch.append(it)
            except:
                pass

        for it in dentists_to_fetch:
            d_url = it.get('url')
            d_name = it.get('name')
            try:
                r_det = requests.get(d_url, headers=HEADERS, timeout=8)
                if r_det.status_code != 200:
                    continue
                s_det = BeautifulSoup(r_det.text, 'html.parser')
                tel = ""
                addr_text = ""
                for s in s_det.find_all('script', type='application/ld+json'):
                    try:
                        d = json.loads(s.string)
                        if isinstance(d, dict) and d.get('@type') == 'Dentist':
                            tel = d.get('telephone') or ""
                            ad = d.get('address', {})
                            if isinstance(ad, dict):
                                addr_text = ad.get('streetAddress') or ""
                    except:
                        pass

                wa_tag = s_det.find('a', href=re.compile(r'wa\.me/|whatsapp'))
                wa_href = wa_tag.get('href', '') if wa_tag else ""

                phone_fmt, wa_link = extract_mobile_and_wa(tel, wa_href, s_det.text)
                
                if not phone_fmt or not wa_link:
                    continue

                name_lower = d_name.lower()
                spec = "Odontologia / Cirurgião-Dentista"
                if "orto" in name_lower:
                    spec = "Odontologia - Ortodontia"
                elif "implante" in name_lower:
                    spec = "Odontologia - Implantodontia"
                elif "estética" in name_lower or "faceta" in name_lower:
                    spec = "Odontologia - Estética Dental"
                elif "endodontia" in name_lower or "canal" in name_lower:
                    spec = "Odontologia - Endodontia"
                elif "pediatria" in name_lower or "odontopediatria" in name_lower:
                    spec = "Odontopediatria"

                digits = re.sub(r'\D', '', phone_fmt)
                ddd = digits[:2] if len(digits) >= 2 else ""
                numero = digits[2:] if len(digits) > 2 else ""
                num_fmt = f"{numero[:5]}-{numero[5:]}" if len(numero) == 9 else numero
                bairro, cep, _ = parse_address_parts(addr_text, city_info['nome'], city_info['uf'])

                results.append({
                    'Nome do Profissional / Clínica': d_name.strip(),
                    'Especialidade': spec,
                    'DDD': ddd,
                    'Número': num_fmt,
                    'WhatsApp Completo': f"({ddd}) {num_fmt}",
                    'Link WhatsApp': wa_link,
                    'Cidade': city_info['nome'],
                    'Estado (UF)': city_info['uf'],
                    'Bairro': bairro,
                    'CEP': cep,
                    'Endereço Completo': addr_text.strip() if addr_text else f"{city_info['nome']} - {city_info['uf']}",
                    'Link da Fonte': d_url
                })
            except:
                pass
    except Exception:
        pass
    return results

def save_to_files(records, base_name='profissionais_saude_whatsapp'):
    fieldnames = [
        'Nome do Profissional / Clínica',
        'Especialidade',
        'DDD',
        'Número',
        'WhatsApp Completo',
        'Link WhatsApp',
        'Cidade',
        'Estado (UF)',
        'Bairro',
        'CEP',
        'Endereço Completo',
        'Link da Fonte'
    ]

    clean_base = base_name.replace('.csv', '').replace('.xlsx', '')
    csv_file = f"{clean_base}.csv"
    xlsx_file = f"{clean_base}.xlsx"

    with open(csv_file, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=';')
        writer.writeheader()
        writer.writerows(records)

    print(f"[CSV] Salvo em: {os.path.abspath(csv_file)} (delimitador ';')")

    if HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Contatos WhatsApp"

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        border_side = Side(style="thin", color="D9D9D9")
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")

        ws.append(fieldnames)
        for col_num in range(1, len(fieldnames) + 1):
            c = ws.cell(row=1, column=col_num)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx, r in enumerate(records, start=2):
            row_data = [r[k] for k in fieldnames]
            ws.append(row_data)
            for col_idx in range(1, len(row_data) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font = regular_font
                c.border = cell_border
                if row_idx % 2 == 0:
                    c.fill = zebra_fill
                if col_idx in (3, 4, 5, 8, 10):
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in (6, 12) and c.value:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                    c.font = link_font
                    c.hyperlink = c.value
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 10)

        ws.freeze_panes = "A2"
        wb.save(xlsx_file)
        print(f"[EXCEL] Salvo em: {os.path.abspath(xlsx_file)} (.xlsx nativo)")

def run_extraction(selected_cities, max_pages=3, output_base='profissionais_saude_whatsapp'):
    print("=" * 70)
    print("INICIANDO EXTRAÇÃO DE WHATSAPP COM COLUNAS SEPARADAS")
    print(f"Cidades: {', '.join([c['nome'] for c in selected_cities])}")
    print(f"Páginas por cidade: {max_pages}")
    print("=" * 70)

    all_records = []
    tasks = []
    with ThreadPoolExecutor(max_workers=8) as executor:
        for city in selected_cities:
            for page in range(1, max_pages + 1):
                tasks.append(executor.submit(scrape_dentmap_page, city, page))

        completed = 0
        total = len(tasks)
        for fut in as_completed(tasks):
            all_records.extend(fut.result())
            completed += 1
            print(f"[{completed}/{total}] Páginas processadas... ({len(all_records)} contatos)")

    seen = set()
    unique = []
    for r in all_records:
        digits = re.sub(r'\D', '', r['WhatsApp Completo'])
        if digits and digits not in seen:
            seen.add(digits)
            unique.append(r)

    unique.sort(key=lambda x: (x['Cidade'], x['Nome do Profissional / Clínica']))
    save_to_files(unique, output_base)
    print("=" * 70)
    print(f"[SUCESSO] {len(unique)} contatos únicos salvos em colunas individuais!")
    print("=" * 70)
    return unique

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Extrator de WhatsApp com Colunas Separadas")
    parser.add_argument('--cidades', type=str, default='todas', help="IDs ou 'todas'")
    parser.add_argument('--paginas', type=int, default=3, help="Páginas por cidade")
    parser.add_argument('--saida', type=str, default='profissionais_saude_whatsapp', help="Nome base do arquivo de saída")
    args = parser.parse_args()

    if args.cidades.lower() == 'todas':
        selected = list(CIDADES_DISPONIVEIS.values())
    else:
        ids = [i.strip() for i in args.cidades.split(',') if i.strip() in CIDADES_DISPONIVEIS]
        selected = [CIDADES_DISPONIVEIS[i] for i in ids]
    if not selected:
        selected = list(CIDADES_DISPONIVEIS.values())

    run_extraction(selected, max_pages=args.paginas, output_base=args.saida)
